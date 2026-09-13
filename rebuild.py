# -*- coding: utf-8 -*-
"""Rebuild saudi-akinator.html from akinator-knowledge-base.xlsx.

Run this after editing the workbook:   python rebuild.py
Both files must sit in the same folder as this script.
"""
import json, re, sys, pathlib
from openpyxl import load_workbook

here = pathlib.Path(__file__).parent
xlsx = here / "akinator-knowledge-base.xlsx"
html = here / "saudi-akinator.html"
for f in (xlsx, html):
    if not f.exists():
        sys.exit(f"missing: {f.name}")

wb = load_workbook(xlsx, data_only=True)
qrows = list(wb["Questions"].iter_rows(min_row=2, values_only=True))
QK = [str(r[0]).strip() for r in qrows if r[0]]
QT = [str(r[1]).strip() for r in qrows if r[0]]
ALT = [[str(r[1]).strip()] + [str(x).strip() for x in r[2:4] if x and str(x).strip()]
       for r in qrows if r[0]]

items, problems = [], []
for n, row in enumerate(wb["Items"].iter_rows(min_row=2, values_only=True), start=2):
    if not row or not row[0]:
        continue
    ar, da, cat = [str(x or "").strip() for x in row[:3]]
    en = str(row[4] or "").strip()
    desc = str(row[5] or "").strip()
    try:
        pop = int(row[3])
        if pop not in (1, 2, 3):
            raise ValueError
    except (TypeError, ValueError):
        problems.append(f"row {n} ({ar}): الشهرة is {row[3]!r} -> treated as 2")
        pop = 2
    vals = row[6:6 + len(QK)]
    m = ""
    for i, v in enumerate(vals):
        v = str(v).strip() if v is not None else "?"
        if v not in ("0", "1", "?"):
            problems.append(f"row {n} ({ar}) column '{QK[i]}': {v!r} -> treated as ?")
            v = "?"
        m += v
    if len(m) != len(QK):
        problems.append(f"row {n} ({ar}): {len(m)} attributes, expected {len(QK)} -> padded")
        m = m.ljust(len(QK), "?")
    items.append({"ar": ar, "en": en, "cat": cat, "d": desc, "da": da, "p": pop, "m": m})

if problems:
    print("warnings:")
    for p in problems[:25]:
        print("  -", p)
    if len(problems) > 25:
        print(f"  ... and {len(problems)-25} more")

data = json.dumps({"q": QT, "alt": ALT, "qk": QK, "items": items},
                  ensure_ascii=False, separators=(",", ":"))
src = html.read_text(encoding="utf-8")
new, n = re.subn(r"const D = \{.*?\};\nconst N", "const D = " + data + ";\nconst N",
                 src, count=1, flags=re.S)
if n != 1:
    sys.exit("could not locate the data block in saudi-akinator.html")
html.write_text(new, encoding="utf-8")
print(f"rebuilt saudi-akinator.html with {len(items)} items and {len(QK)} questions")
